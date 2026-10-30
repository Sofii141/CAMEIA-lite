"""Aplica el estilo "Macarons" a hojas de Excel (openpyxl).

Uso como script: python _Estilo/estilo_xlsx.py "ruta/archivo.xlsx"
  → aplica el estilo a todas las hojas, con el encabezado en la fila 1, y guarda el mismo archivo.
     (Para no sobrescribir una versión entregada, copien antes el archivo con fecha y versión nuevas.)
Uso como módulo: from estilo_xlsx import estilizar_hoja
"""
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MATCHA, SEMOLINA, STRAWBERRY, TEA_ROSE, PEANUT, CHOCOLATE = "BCC86F", "F5CB99", "FFABA9", "F68578", "B9794A", "79371D"
CREMA, SEMOLINA_SOFT, LINEA, TINTA = "FFF8EF", "FCEEDC", "E9CFAE", "4A2312"
# Las fuentes del proyecto (Nunito Sans, Fraunces) no vienen en Windows: en Excel se usa Segoe UI.
FUENTE = "Segoe UI"

borde = Border(bottom=Side(style="thin", color=LINEA))


def relleno(color):
    return PatternFill("solid", fgColor=color)


def estilizar_hoja(ws, fila_encabezado=1, anchos=None, color_pestana=TEA_ROSE):
    """Encabezado chocolate con texto crema, filas alternas semolina, bordes suaves y fila congelada."""
    ws.sheet_properties.tabColor = color_pestana
    for celda in ws[fila_encabezado]:
        celda.fill = relleno(CHOCOLATE)
        celda.font = Font(name=FUENTE, bold=True, color=CREMA, size=10)
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[fila_encabezado].height = 42
    for fila in ws.iter_rows(min_row=fila_encabezado + 1, max_row=max(ws.max_row, fila_encabezado + 1)):
        par = (fila[0].row - fila_encabezado) % 2 == 0
        for celda in fila:
            celda.font = Font(name=FUENTE, size=10, color=TINTA)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            celda.border = borde
            if par:
                celda.fill = relleno(SEMOLINA_SOFT)
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)
    for i, ancho in enumerate(anchos or [], start=1):
        ws.column_dimensions[ws.cell(row=fila_encabezado, column=i).column_letter].width = ancho


if __name__ == "__main__":
    libro = load_workbook(sys.argv[1])
    for hoja in libro:
        estilizar_hoja(hoja)
    libro.save(sys.argv[1])
    print(f"Estilo aplicado: {sys.argv[1]}")
