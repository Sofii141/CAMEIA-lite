"""Resume la hoja "Medicion por caso" de la bitácora: mediana y rango por condición y tamaño.

Uso:
  python "07 - Analisis/Scripts/resumen_medicion.py" "05 - Ciclos y bitacora/29102026_01_Bitacora_v2.xlsx"
  python "07 - Analisis/Scripts/resumen_medicion.py" <bitacora.xlsx> --html "09 - Control de entrega/29102026_02_Punto_Control_v2.html"
      → además reemplaza la sección de línea base del HTML (entre <!-- LB:INICIO --> y <!-- LB:FIN -->) y regenera el PDF.

Calcula todo desde los datos crudos (inicio, fin, pausas…), no desde las fórmulas del Excel: así el resumen de la hoja
y el de este script son dos cálculos independientes. Igual, una persona recalcula a mano al menos una cifra (regla de la E1).
Por qué mediana y rango, y no promedio ni prueba de hipótesis: ver 01 - Exposicion/Fuentes/29102026_02_Como_Medir_y_Presentar_v2.md
"""
import datetime as dt
import html
import re
import statistics
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
RAIZ = Path(__file__).resolve().parents[2]
CONDICIONES = ("Línea base", "Ciclo 1", "Ciclo 2")
TAMANOS = ("S", "M", "L")


def minutos(v):
    """Convierte una hora de Excel (time, datetime, fracción de día o texto hh:mm) a minutos desde medianoche."""
    if v is None or v == "":
        return None
    if isinstance(v, (dt.time, dt.datetime)):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)):
        return round(v * 1440)
    m = re.match(r"\s*(\d{1,2}):(\d{2})", str(v))
    return int(m[1]) * 60 + int(m[2]) if m else None


def numero(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def leer(xlsx):
    hoja = load_workbook(xlsx, data_only=True)["Medicion por caso"]
    filas = []
    for r in hoja.iter_rows(min_row=2, values_only=True):
        caso, tarea, tam, dif, cond = r[0], r[1], r[2], r[3], r[4]
        ini, fin = minutos(r[8]), minutos(r[9])
        if not cond or ini is None or fin is None or (r[22] or "").strip() == "[C]":
            continue
        dur = fin - ini - (numero(r[10]) or 0)
        personas = 2 if cond == "Línea base" else 1
        filas.append(dict(caso=caso, tarea=tarea, tam=tam or "?", dif=dif == "Sí", cond=cond, quien=r[5] or "", pron=numero(r[7]),
                          dur=dur, esf=dur * personas, falla=numero(r[14]), total=numero(r[15]), retr=numero(r[16]),
                          cw_a=numero(r[18]), cw_c=numero(r[19]), ayuda=numero(r[20]), marca=r[22] or ""))
    return filas


def leer_sesiones(xlsx, etapa="Línea base"):
    hoja = load_workbook(xlsx, data_only=True)["Sesiones"]
    cab = [c.value for c in hoja[1]]
    filas = []
    for r in hoja.iter_rows(min_row=2, values_only=True):
        if r[2] == etapa and r[1]:  # solo sesiones que ya ocurrieron (con fecha)
            filas.append({k: ("" if v is None else v.strftime("%H:%M") if isinstance(v, (dt.time, dt.datetime)) else str(v))
                          for k, v in zip(cab, r)})
    return filas


def med(valores):
    v = [x for x in valores if x is not None]
    if not v:
        return "—"
    m = statistics.median(v)
    txt = f"{m:g}" if float(m).is_integer() else f"{m:.1f}"
    return txt if len(v) == 1 else f"{txt} [{min(v):g}–{max(v):g}]"


def a12(x, y):
    """Vargha y Delaney (2000): probabilidad de que un valor de x sea mayor que uno de y (0,5 = sin diferencia)."""
    if not x or not y:
        return None
    return sum((a > b) + 0.5 * (a == b) for a in x for b in y) / (len(x) * len(y))


def resumen(filas):
    salida = []
    for cond in CONDICIONES:
        fc = [f for f in filas if f["cond"] == cond]
        if not fc:
            continue
        for tam in TAMANOS + ("Total",):
            g = fc if tam == "Total" else [f for f in fc if f["tam"] == tam]
            if not g:
                continue
            falla = sum(f["falla"] or 0 for f in g)
            total = sum(f["total"] or 0 for f in g)
            ratio = [f["dur"] / f["pron"] for f in g if f["pron"]]
            salida.append(dict(cond=cond, tam=tam, n=len(g), dur=med(f["dur"] for f in g), esf=med(f["esf"] for f in g),
                               ocultas=f"{falla:g}/{total:g} ({100 * falla / total:.0f} %)" if total else "—",
                               retr=med(f["retr"] for f in g), cw_a=med(f["cw_a"] for f in g), cw_c=med(f["cw_c"] for f in g),
                               ratio=med(round(x, 2) for x in ratio)))
    return salida


def imprimir(filas, res):
    print(f"Tareas con datos válidos: {len(filas)}\n")
    cab = ("Condición", "Tam.", "n", "Duración", "Esfuerzo", "Ocultas que fallan", "Retrabajo", "Codewalk autor", "Codewalk comp.", "Real/pronóstico")
    print(" | ".join(cab))
    for r in res:
        print(" | ".join(str(r[k]) for k in ("cond", "tam", "n", "dur", "esf", "ocultas", "retr", "cw_a", "cw_c", "ratio")))
    base = [f["dur"] for f in filas if f["cond"] == "Línea base"]
    agente = [f["dur"] for f in filas if f["cond"] in ("Ciclo 1", "Ciclo 2")]
    if base and agente:
        print(f"\nA12 duración (agente > línea base): {a12(agente, base):.2f}  ·  0,5 = igual; descriptivo, no prueba causalidad.")
    print("\nValores = mediana [mínimo–máximo]. Las tareas marcadas [C] no se cuentan.")


def fragmento_html(filas, res, sesiones):
    e = html.escape
    base = [f for f in filas if f["cond"] == "Línea base"]
    # Fecha, horas, participantes, lugar, herramientas, IA, semilla e interrupciones (las condiciones que pide el enunciado)
    claves = [k for i, k in enumerate(list(sesiones[0]) if sesiones else []) if i in (1, 3, 4, 5, 6, 7, 8, 9, 11)]
    sesiones_html = "".join("<tr>" + "".join(f"<td>{e(s[k])}</td>" for k in claves) + "</tr>" for s in sesiones)
    cab_sesiones = "".join(f"<th>{e(k)}</th>" for k in claves)
    condiciones = (f"<h3>Condiciones de la medición</h3><table class='pequeno'><thead><tr>{cab_sesiones}</tr></thead><tbody>{sesiones_html}</tbody></table>"
                   if sesiones else "<div class='caja alerta'><h4>Faltan las condiciones</h4><p>La hoja Sesiones no tiene filas de línea base.</p></div>")
    def v(x):
        return "—" if x is None else f"{x:g}"

    dif = "<span class='chip dif'>Sí</span>"
    filas_html = "".join(
        f"<tr><td><code>{e(str(f['caso']))}</code> {e(str(f['tarea']))}</td><td class='c'><span class='chip {f['tam'].lower()}'>{e(f['tam'])}</span></td>"
        f"<td class='c'>{dif if f['dif'] else ''}</td><td class='c'>{v(f['pron'])}</td><td class='c'>{v(f['dur'])}</td><td class='c'>{v(f['esf'])}</td>"
        f"<td class='c'>{v(f['falla'] or 0)}/{v(f['total'] or 0)}</td><td class='c'>{v(f['retr'])}</td>"
        f"<td class='c'>{v(f['cw_a'])} / {v(f['cw_c'])}</td></tr>"
        for f in base)
    res_html = "".join(
        f"<tr><td>{e(r['tam'])}</td><td class='c'>{r['n']}</td><td>{r['dur']}</td><td>{r['esf']}</td><td>{r['ocultas']}</td>"
        f"<td>{r['retr']}</td><td>{r['cw_a']} / {r['cw_c']}</td><td>{r['ratio']}</td></tr>"
        for r in res if r["cond"] == "Línea base")
    if not base:
        return "<div class='caja alerta'><h4>Sin datos todavía</h4><p>La bitácora aún no tiene tareas de línea base con inicio y fin.</p></div>"
    return (condiciones + f"<h3>Resultados por tarea (n = {len(base)})</h3><table><thead><tr><th>Tarea</th><th>Tam.</th><th>Difícil</th><th>Pronóstico</th>"
            "<th>Duración (min)</th><th>Esfuerzo (min·pers.)</th><th>Ocultas que fallan</th><th>Retrabajo (min)</th><th>Codewalk autor / comp.</th></tr></thead>"
            f"<tbody>{filas_html}</tbody></table>"
            "<h3>Resumen por tamaño</h3><p class='pequeno'>Mediana [mínimo–máximo]. Real/pronóstico &gt; 1 = tardaron más de lo que creían.</p>"
            "<table><thead><tr><th>Tamaño</th><th>n</th><th>Duración</th><th>Esfuerzo</th><th>Ocultas que fallan</th><th>Retrabajo</th>"
            f"<th>Codewalk autor / comp.</th><th>Real/pronóstico</th></tr></thead><tbody>{res_html}</tbody></table>")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    filas = leer(sys.argv[1])
    res = resumen(filas)
    imprimir(filas, res)
    if "--html" in sys.argv:
        destino = Path(sys.argv[sys.argv.index("--html") + 1])
        texto = destino.read_text(encoding="utf-8")
        nuevo = re.sub(r"(<!-- LB:INICIO -->).*?(<!-- LB:FIN -->)", lambda m: m[1] + fragmento_html(filas, res, leer_sesiones(sys.argv[1])) + m[2], texto, flags=re.S)
        destino.write_text(nuevo, encoding="utf-8")
        subprocess.run([sys.executable, str(RAIZ / "_Estilo" / "html_a_pdf.py"), str(destino)], check=True)
